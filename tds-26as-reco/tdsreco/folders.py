"""Creates the working folder and its input templates (`python -m tdsreco init <folder>`)."""
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .config import AREA_KEYWORDS, DEFAULT_SECTION_MAP, DEFAULT_STANDARD_RATES, DEFAULTS
from .ldc import LDC_COLS
from .masters import PARTY_COLS, TAN_COLS
from .tracker import STATUS_CHOICES

FOLDERS = {
    "00_Config": "Settings.xlsx - tolerance, cost of capital, cut-off override, GL accounts, section map, owners.",
    "01_26AS": "Drop every Form 26AS you download from TRACES (text .txt or Excel), any FY, any date. Keep the "
               "old ones - the newest per FY is used and the older ones give 'what changed'. Tip: put the "
               "download date in the file name, e.g. 26AS_FY2025-26_2026-09-15.txt",
    "02_Books_GL": "Drop TDS/TCS receivable GL line-item dumps (SAP FBL3N/FAGLL03 export). Full ITD dumps or "
                   "monthly increments - overlaps are de-duplicated.",
    "03_Masters": "Party_Master.xlsx (TAN -> PAN, area, customer codes, e-mail) and LDC_Certificates.xlsx.",
    "04_Party_Ledgers": "Ledgers received from parties (their books, our account). Put the party PAN or TAN in "
                        "the file name: AAJCG0980D_ledger_FY25-26.xlsx",
    "05_Form16A": "Form 16A PDFs or TRACES zips received from deductors. Scanned PDFs: key them into "
                  "Form16A_Manual_Entry.xlsx instead.",
    "06_AR_Open_Items": "Optional. Open customer invoices with expected TDS - explains 'short in books' that is "
                        "only TDS on invoices not yet paid.",
    "90_Output": "Reco workbooks written by each run.",
    "91_Tracker": "Action_Tracker.xlsx - the ONE file people edit: Status, Remarks, Assigned To, Next Follow-up.",
    "99_History": "Run snapshots (used for 'movement vs previous run'). Do not edit.",
}


def _head(ws, fill="1F3864"):
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=fill)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, min(60, len(str(col[0].value or "")) + 4))


def init_folder(root, seed_gl_accounts=None, seed_owners=None):
    os.makedirs(root, exist_ok=True)
    for f, desc in FOLDERS.items():
        d = os.path.join(root, f)
        os.makedirs(d, exist_ok=True)
        with open(os.path.join(d, "_README.txt"), "w") as fh:
            fh.write(desc + "\n")

    s = os.path.join(root, "00_Config", "Settings.xlsx")
    if not os.path.exists(s):
        with pd.ExcelWriter(s, engine="openpyxl") as w:
            pd.DataFrame([(k, v[0], v[1]) for k, v in DEFAULTS.items()], columns=["Key", "Value", "Description"]) \
                .to_excel(w, sheet_name="Settings", index=False)
            (seed_gl_accounts if seed_gl_accounts is not None else
             pd.DataFrame(columns=["GL Code", "Description", "FY", "Area", "Kind", "Use for JE"])) \
                .to_excel(w, sheet_name="GL_Accounts", index=False)
            pd.DataFrame(DEFAULT_SECTION_MAP, columns=["26AS Section", "Section"]) \
                .to_excel(w, sheet_name="Section_Map", index=False)
            pd.DataFrame(DEFAULT_STANDARD_RATES, columns=["Section", "Standard Rate", "Effective From", "Effective To"]) \
                .to_excel(w, sheet_name="Standard_Rates", index=False)
            (seed_owners if seed_owners is not None else
             pd.DataFrame([(a, f"{a} team", "") for a in sorted({v for _, v in AREA_KEYWORDS} |
                                                              {"Sellers", "Inter company", "AP Trade", "BFD & Monet"})],
                          columns=["Area", "Owner", "Email"])).to_excel(w, sheet_name="Owners", index=False)
        wb = load_workbook(s)
        for ws in wb.worksheets:
            _head(ws)
        ws = wb["Settings"]
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 24, 14, 110
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="FFF2CC")
            ws.cell(row=r, column=2).font = Font(name="Arial", color="0000FF", size=9)
        wb.save(s)

    pm = os.path.join(root, "03_Masters", "Party_Master.xlsx")
    if not os.path.exists(pm):
        with pd.ExcelWriter(pm, engine="openpyxl") as w:
            pd.DataFrame(columns=TAN_COLS).to_excel(w, sheet_name="TAN_Map", index=False)
            pd.DataFrame(columns=PARTY_COLS).to_excel(w, sheet_name="Parties", index=False)
        wb = load_workbook(pm)
        for ws in wb.worksheets:
            _head(ws)
        wb.save(pm)

    lp = os.path.join(root, "03_Masters", "LDC_Certificates.xlsx")
    if not os.path.exists(lp):
        with pd.ExcelWriter(lp, engine="openpyxl") as w:
            pd.DataFrame(columns=LDC_COLS).to_excel(w, sheet_name="LDC", index=False)
            pd.DataFrame({"How to fill": [
                "One row per certificate (u/s 197 / its new-Act equivalent), exactly as on the certificate.",
                "Deductor TAN: the TAN the certificate is addressed to. Use ALL only for a certificate valid for "
                "any deductor.",
                "Section: the section as the reco sees it (194C, 194JB, 194Q ...). New-Act codes in 26AS are "
                "mapped via Settings > Section_Map.",
                "LDC Rate (%): the percentage, e.g. 0.25 for 0.25%.",
                "Valid From / Valid To: dates. Amount Limit (Rs): the amount on the certificate; blank = no limit.",
                "Example: LDC123456789 | MUMA12345B | ABC Ltd | 194C | 0.25 | 01-04-2025 | 31-03-2026 | 500000000",
                "Tip: the run's 'LDC Master Suggestions' sheet lists rates seen in 26AS that look like an LDC.",
            ]}).to_excel(w, sheet_name="Instructions", index=False)
        wb = load_workbook(lp)
        _head(wb["LDC"])
        wb["Instructions"].column_dimensions["A"].width = 120
        wb.save(lp)

    fp = os.path.join(root, "05_Form16A", "Form16A_Manual_Entry.xlsx")
    if not os.path.exists(fp):
        with pd.ExcelWriter(fp, engine="openpyxl") as w:
            pd.DataFrame(columns=["TAN", "Deductor Name", "FY", "Quarter", "Certificate No", "Amount Paid",
                                  "Tax Deducted", "Tax Deposited"]).to_excel(w, sheet_name="Form16A", index=False)
            pd.DataFrame({"How to fill": [
                "Use only for certificates the tool cannot read (scans). One row per TAN x quarter.",
                "FY as 2025-26, Quarter as Q1..Q4. Example: MUMA12345B | ABC Ltd | 2025-26 | Q2 | ABCDEFG | "
                "1000000 | 2500 | 2500"]}).to_excel(w, sheet_name="Instructions", index=False)
        wb = load_workbook(fp)
        _head(wb["Form16A"])
        wb.save(fp)

    lt = os.path.join(root, "04_Party_Ledgers", "_Template_Party_Ledger.xlsx")
    if not os.path.exists(lt):
        pd.DataFrame(columns=["Date", "Document No", "Particulars", "Debit", "Credit", "TDS Amount", "Section"]) \
            .to_excel(lt, index=False)

    ar = os.path.join(root, "06_AR_Open_Items", "_Template_AR_Open_Items.xlsx")
    if not os.path.exists(ar):
        pd.DataFrame(columns=["PAN", "Customer Code", "Invoice No", "Invoice Date", "Taxable Value", "Expected TDS"]) \
            .to_excel(ar, index=False)

    tp = os.path.join(root, "91_Tracker", "Action_Tracker.xlsx")
    if not os.path.exists(tp):
        from .tracker import SYS_COLS, USER_COLS
        pd.DataFrame(columns=SYS_COLS + USER_COLS).to_excel(tp, sheet_name="Tracker", index=False)
    return root


def style_tracker(path):
    wb = load_workbook(path)
    ws = wb["Tracker"]
    _head(ws)
    hdr = [c.value for c in ws[1]]
    user_fill = PatternFill("solid", fgColor="FFF2CC")
    for name in ("Status", "Remarks", "Assigned To", "Next Follow-up", "Updated By"):
        if name in hdr:
            j = hdr.index(name) + 1
            ws.cell(row=1, column=j).fill = PatternFill("solid", fgColor="BF8F00")
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=j).fill = user_fill
    if "Status" in hdr:
        col = ws.cell(row=1, column=hdr.index("Status") + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_CHOICES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{max(ws.max_row, 2) + 2000}")
    for name, w_ in (("Action", 70), ("Remarks", 40), ("Party", 34), ("Category", 34)):
        if name in hdr:
            ws.column_dimensions[ws.cell(row=1, column=hdr.index(name) + 1).column_letter].width = w_
    for name in ("Amount (this run)", "Amount (first seen)"):
        if name in hdr:
            j = hdr.index(name) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=j).number_format = '#,##0;(#,##0);"-"'
    for name in ("First Seen", "Last Seen", "Next Follow-up"):
        if name in hdr:
            j = hdr.index(name) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=j).number_format = "dd-mmm-yyyy"
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


# ---------------------------------------------------------------- portable bundle
LAUNCH_BAT = r"""@echo off
REM Double-click: reconcile everything in this folder.
setlocal
cd /d "%~dp0"
set PYTHONPATH=%~dp0_engine
python -c "import pandas, openpyxl" 2>nul || python -m pip install -r "_engine\requirements.txt"
python -m tdsreco {cmd} "%~dp0."
pause
"""

LAUNCH_SH = """#!/bin/bash
# Double-click (Mac) or run: reconcile everything in this folder.
cd "$(dirname "$0")"
export PYTHONPATH="$PWD/_engine"
python3 -c "import pandas, openpyxl" 2>/dev/null || python3 -m pip install --user -r _engine/requirements.txt
python3 -m tdsreco {cmd} "$PWD"
read -r -p "Done. Press Enter to close."
"""

START_HERE = """26AS vs BOOKS - AUTOMATED RECONCILIATION
=========================================

First time (once)
  1. Install Python 3.10+ (python.org; on Windows tick "Add python.exe to PATH").
  2. Fill 00_Config/Settings.xlsx (yellow cells): company name, company PAN, cost of capital.
  3. Fill 03_Masters/LDC_Certificates.xlsx with your lower deduction certificates.

Every time
  1. Drop new files into the numbered folders (see _README.txt in each):
       01_26AS            26AS downloads (any FY, keep old ones)
       02_Books_GL        TDS receivable GL dumps
       04_Party_Ledgers   ledgers received from parties (PAN/TAN in file name)
       05_Form16A         Form 16A PDFs / zips
  2. Double-click RUN_RECO.bat (Windows) or RUN_RECO.command (Mac).
     WATCH_FOLDER.bat / .command keeps running and re-reconciles whenever a file lands.
  3. Open the newest workbook in 90_Output (Dashboard, then Action Items).
  4. Update status/remarks ONLY in 91_Tracker/Action_Tracker.xlsx - they carry forward.

Full playbook: README.md in this folder.
"""


def bundle(root, package_dir, repo_dir):
    """Make root a self-contained folder: engine copy + launchers + templates."""
    import shutil
    init_folder(root)
    eng = os.path.join(root, "_engine")
    if os.path.exists(os.path.join(eng, "tdsreco")):
        shutil.rmtree(os.path.join(eng, "tdsreco"))
    shutil.copytree(package_dir, os.path.join(eng, "tdsreco"), ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
    for f in ("requirements.txt",):
        shutil.copy2(os.path.join(repo_dir, f), os.path.join(eng, f))
    shutil.copy2(os.path.join(repo_dir, "README.md"), os.path.join(root, "README.md"))
    for name, cmd in (("RUN_RECO", "run"), ("WATCH_FOLDER", "watch")):
        with open(os.path.join(root, f"{name}.bat"), "w", newline="\r\n") as fh:
            fh.write(LAUNCH_BAT.format(cmd=cmd))
        p = os.path.join(root, f"{name}.command")
        with open(p, "w", newline="\n") as fh:
            fh.write(LAUNCH_SH.format(cmd=cmd))
        os.chmod(p, 0o755)
    with open(os.path.join(root, "START_HERE.txt"), "w") as fh:
        fh.write(START_HERE)
    return root
