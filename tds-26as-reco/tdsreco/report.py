"""Writes 90_Output/26AS_Reco_<cutoff>_run<date>.xlsx."""
import datetime as dt

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .reco import WHO_DED, WHO_NONE, WHO_OK, WHO_TAG, WHO_TAX, WHO_US
from .util import clean, fy_label

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
NUM = '#,##0;(#,##0);"-"'
PCT = '0.00%;(0.00%);"-"'
DATE = "dd-mmm-yyyy"


def _style_table(ws, df, start_row, money=(), pct=(), dates=(), widths=None, input_cols=()):
    hdr = start_row
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=hdr, column=j)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        cell.fill = HDR_FILL if c not in input_cols else PatternFill("solid", fgColor="BF8F00")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    n = len(df)
    col_idx = {c: j for j, c in enumerate(df.columns, 1)}
    fmts = {**{c: NUM for c in money}, **{c: PCT for c in pct}, **{c: DATE for c in dates}}
    body = Font(name=FONT, size=9)
    for r in ws.iter_rows(min_row=hdr + 1, max_row=hdr + n, max_col=len(df.columns)):
        for cell in r:
            cell.font = body
    for c, f in fmts.items():
        if c in col_idx:
            j = col_idx[c]
            for i in range(hdr + 1, hdr + n + 1):
                ws.cell(row=i, column=j).number_format = f
    for c in input_cols:
        if c in col_idx:
            j = col_idx[c]
            for i in range(hdr + 1, hdr + n + 1):
                ws.cell(row=i, column=j).fill = INPUT_FILL
    for j, c in enumerate(df.columns, 1):
        w = (widths or {}).get(c)
        if w is None:
            sample = [len(clean(c))] + [len(clean(v)) for v in df[c].head(200).tolist()]
            w = min(max(sample) + 2, 60)
            if c in fmts:
                w = max(12, min(w, 16))
        ws.column_dimensions[get_column_letter(j)].width = max(w, 8)
    ws.row_dimensions[hdr].height = 30
    if n:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(df.columns))}{hdr + n}"
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)


def _write(writer, name, df, title=None, note=None, money=(), pct=(), dates=(), widths=None, total_cols=(),
           input_cols=()):
    df = df.copy()
    for c in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            df[c] = df[c].dt.tz_localize(None) if getattr(df[c].dt, "tz", None) else df[c]
    start = 4
    df.to_excel(writer, sheet_name=name, index=False, startrow=start - 1)
    ws = writer.sheets[name]
    ws["A1"] = title or name
    ws["A1"].font = Font(name=FONT, bold=True, size=13, color="1F3864")
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(name=FONT, italic=True, size=9, color="595959")
    _style_table(ws, df, start, money, pct, dates, widths, input_cols)
    # subtotal row above the header, like the manual file (respects filters)
    if len(df):
        for c in total_cols:
            if c in df.columns:
                j = list(df.columns).index(c) + 1
                col = get_column_letter(j)
                cell = ws.cell(row=start - 1, column=j, value=f"=SUBTOTAL(9,{col}{start + 1}:{col}{start + len(df)})")
                cell.number_format = NUM
                cell.font = Font(name=FONT, bold=True, size=9)
                cell.fill = SUB_FILL
    return ws


def _block(ws, row, title, df, money=(), pct=()):
    ws.cell(row=row, column=1, value=title).font = Font(name=FONT, bold=True, size=11, color="1F3864")
    row += 1
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=row, column=j, value=str(c))
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True)
    for i, rec in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(rec, 1):
            if isinstance(v, (np.floating, float)) and np.isnan(v):
                v = None
            if isinstance(v, np.generic):
                v = v.item()
            cell = ws.cell(row=row + i, column=j, value=v)
            cell.font = Font(name=FONT, size=9, bold=str(rec[0]).startswith("Total"))
            c = df.columns[j - 1]
            if c in money:
                cell.number_format = NUM
            elif c in pct:
                cell.number_format = PCT
            cell.border = Border(bottom=THIN)
    return row + len(df) + 3


def email_draft(cfg, party_name, key, items, cutoff):
    comp = cfg["company_name"] or "our company"
    pan = cfg["company_pan"] or "<our PAN>"
    lines = [f"Subject: TDS credit mismatch in Form 26AS - {comp} (PAN {pan})", "",
             f"Dear {party_name or 'Sir/Madam'} team,", "",
             f"While reconciling Form 26AS of {comp} (PAN {pan}) up to {cutoff:%d-%b-%Y} with our books, we found "
             "the following points against your TAN. Please help us close them:", ""]
    for i, it in enumerate(items, 1):
        lines.append(f"{i}. {it}")
    lines += ["", "Please share your ledger of our account and the Form 16A for the periods above, and file a "
                  "correction statement where required. Happy to get on a call.", "", "Regards,"]
    return "\n".join(lines)


def write_report(path, ctx):
    cfg, rc = ctx["cfg"], ctx["reco"]
    keys = rc.keys.copy()
    cutoff = rc.cutoff
    tol = cfg["tolerance"]
    fys = rc.fys

    # --------------------------------------------------------------- reco sheet
    fyv = rc.fyv
    wide = keys[["party_key"]].copy()
    for f in fys:
        sub = fyv[fyv["fy"] == f].set_index("party_key")
        wide[f"{fy_label(f)} Books"] = wide["party_key"].map(sub["books"]).fillna(0.0)
        wide[f"{fy_label(f)} 26AS"] = wide["party_key"].map(sub["s26"]).fillna(0.0)
        wide[f"{fy_label(f)} Variance"] = wide[f"{fy_label(f)} Books"] - wide[f"{fy_label(f)} 26AS"]
    trk = ctx["tracker"].set_index("Action ID") if len(ctx["tracker"]) else pd.DataFrame()
    led = ctx["ledger_info"]
    last_led = {}
    if len(led):
        ok = led[led["party_key"] != ""]
        last_led = ok.groupby("party_key")["received_on"].max().to_dict()
    reco = pd.DataFrame({
        "PAN/TAN": keys["party_key"], "Name": keys["name"], "Name as per 26AS": keys["name_26as"],
        "TAN(s)": keys["tans"], "Sections": keys["sections"], "Area": keys["area"], "Owner": keys["owner"],
        "Last party ledger received": keys["party_key"].map(last_led),
    })
    reco = pd.concat([reco, wide.drop(columns=["party_key"])], axis=1)
    reco["Books ITD"] = keys["books"]
    reco["26AS ITD"] = keys["s26"]
    reco["Variance ITD"] = keys["variance"]
    reco["of which 26AS U/O"] = keys["s26_uo"]
    reco["Booked after cut-off (same period)"] = keys["sub_books_pre"]
    reco["Booked after cut-off (later/unknown period)"] = keys["sub_books_other"]
    reco["26AS after cut-off"] = keys["sub_26as"]
    reco["Delta after subsequent"] = keys["delta_after_sub"]
    reco["Party ledger TDS"] = keys["ledger_tds"]
    reco["Form 16A TDS"] = keys["form16a_tds"]
    reco["Excess/ Short"] = keys["status"]
    reco["Category"] = keys["category"]
    reco["Who acts"] = keys["who"]
    reco["Action"] = keys["action"]
    reco["Priority"] = keys["priority"]
    reco["Action ID"] = keys["action_id"]
    if len(trk):
        reco["Tracker status"] = keys["action_id"].map(trk["Status"]).fillna("")
        reco["Tracker remarks"] = keys["action_id"].map(trk["Remarks"]).fillna("")
    prev = ctx.get("prev_keys")
    if prev is not None:
        pv = prev.set_index("party_key")
        reco["Previous run variance"] = keys["party_key"].map(pv["variance"]).fillna(0.0)
        reco["Change vs previous"] = reco["Variance ITD"] - reco["Previous run variance"]
        reco["Previous status"] = keys["party_key"].map(pv["status"]).fillna("")
    money_reco = [c for c in reco.columns if any(w in c for w in ("Books", "26AS", "Variance", "cut-off",
                                                                   "Delta", "TDS", "U/O", "Change"))
                  and c not in ("Name as per 26AS",)]

    # --------------------------------------------------------------- actions
    act = keys[~keys["who"].isin([WHO_OK, WHO_NONE])].copy()
    fy_txt = {k: ", ".join(f"{fy_label(f)} {v:,.0f}" for f, v in sorted(rc._fy_split(k).items()) if abs(v) > 1)
              for k in act["party_key"]}
    je_ids = set(rc.journal["Action ID"]) if len(rc.journal) else set()
    actions = pd.DataFrame({
        "Action ID": act["action_id"], "Priority": act["priority"], "Who Acts": act["who"], "Owner": act["owner"],
        "Area": act["area"], "PAN/TAN": act["party_key"], "Party": act["name"], "Category": act["category"],
        "Action": act["action"],
        "Amount": [rc._action_amount(r) * (1 if r["family"] in ("UNALLOC", "FYRECLASS") else np.sign(r["variance"]))
                   if r["family"] != "UNALLOC" else r["books"] for r in act.to_dict("records")],
        "FY split (books - 26AS)": act["party_key"].map(fy_txt), "Evidence": act["evidence"],
        "Proposed JE": act["action_id"].map(lambda a: f"JE-{a}" if a in je_ids else ""),
    })
    extra = []
    uo = rc.extra_actions
    for r in uo.to_dict("records") if len(uo) else []:
        extra.append({"Action ID": f"UO-{r['tan']}-{r['booking_status']}", "Priority": "Medium" if abs(r["tds"]) >= cfg["priority_medium"] else "Low",
                      "Who Acts": WHO_DED, "Owner": cfg.owner_for(keys.set_index("party_key")["area"].get(r["party_key"], "")),
                      "Area": keys.set_index("party_key")["area"].get(r["party_key"], ""),
                      "PAN/TAN": r["party_key"], "Party": r["deductor_name"],
                      "Category": f"26AS status {r['booking_status']} ({r['fys']})", "Action": r["action"],
                      "Amount": r["tds"], "FY split (books - 26AS)": "", "Evidence": f"{r['lines']} lines",
                      "Proposed JE": ""})
    ldc_sum = ctx["ldc_summary"]
    if len(ldc_sum):
        g = ldc_sum[ldc_sum["excess"] > 1].groupby(["party_key", "deductor_name", "tan"]).agg(
            excess=("excess", "sum"), wc=("wc_cost_net", "sum"), certs=("cert", lambda s: ", ".join(sorted(set(s)))),
            secs=("section", lambda s: ", ".join(sorted(set(s))))).reset_index()
        area_of = keys.set_index("party_key")["area"]
        for r in g.to_dict("records"):
            extra.append({"Action ID": f"LDC-{r['tan']}", "Priority": "High" if r["excess"] >= cfg["priority_high"]
                          else "Medium" if r["excess"] >= cfg["priority_medium"] else "Low",
                          "Who Acts": WHO_DED, "Owner": cfg.owner_for(area_of.get(r["party_key"], "")),
                          "Area": area_of.get(r["party_key"], ""), "PAN/TAN": r["party_key"],
                          "Party": r["deductor_name"], "Category": "LDC not applied (working-capital loss)",
                          "Action": f"Send LDC certificate(s) {r['certs']} ({r['secs']}) again and get the LDC rate "
                                    f"applied to all future payments. Net cost so far ~Rs {r['wc']:,.0f}.",
                          "Amount": r["excess"], "FY split (books - 26AS)": "", "Evidence": "LDC Cut",
                          "Proposed JE": ""})
    util = ctx["ldc_util"]
    if len(util):
        for r in util[util["Alert"] != ""].to_dict("records"):
            extra.append({"Action ID": f"LDCCERT-{r['Certificate No']}", "Priority": "High", "Who Acts": WHO_TAX,
                          "Owner": "Tax team", "Area": "", "PAN/TAN": r["Deductor TAN"], "Party": "",
                          "Category": "LDC certificate", "Action": r["Alert"], "Amount": r["Amount Consumed (Rs)"],
                          "FY split (books - 26AS)": "", "Evidence": f"{r['Section']} @ {r['LDC Rate (%)']}%",
                          "Proposed JE": ""})
    if extra:
        actions = pd.concat([actions, pd.DataFrame(extra)], ignore_index=True)
    order = {"High": 0, "Medium": 1, "Low": 2}
    actions = actions.assign(_p=actions["Priority"].map(order).fillna(3), _a=-actions["Amount"].abs()) \
        .sort_values(["_p", "_a"]).drop(columns=["_p", "_a"])
    if len(trk):
        actions["Tracker status"] = actions["Action ID"].map(trk["Status"]).fillna("New")
        actions["Tracker remarks"] = actions["Action ID"].map(trk["Remarks"]).fillna("")
    ctx["actions_out"] = actions

    # --------------------------------------------------------------- deductor requests
    req = []
    ded = actions[actions["Who Acts"] == WHO_DED]
    for key, grp in ded.groupby("PAN/TAN"):
        items = []
        for r in grp.to_dict("records"):
            if r["Category"].startswith("LDC"):
                items.append(f"Lower deduction certificate not applied: TDS of Rs {r['Amount']:,.0f} was deducted "
                             "above the certificate rate. Please apply the certificate rate on all payments.")
            elif r["Category"].startswith("26AS status U"):
                items.append(f"Rs {r['Amount']:,.0f} shows as 'Unmatched' in 26AS - please correct the challan "
                             "details in a correction statement.")
            elif r["Category"].startswith("26AS status O"):
                items.append(f"Rs {r['Amount']:,.0f} shows as 'Overbooked' in 26AS - please pay the balance or "
                             "correct the statement.")
            else:
                items.append(f"TDS of Rs {abs(r['Amount']):,.0f} deducted by you ({r['FY split (books - 26AS)']}) is "
                             "not reflected in our Form 26AS. Please file/revise your TDS return quoting our PAN.")
        name = grp["Party"].iloc[0]
        req.append({"PAN/TAN": key, "Party": name, "Area": grp["Area"].iloc[0], "Owner": grp["Owner"].iloc[0],
                    "E-mail": rc.party.email_of(key), "Points": len(items), "Amount": grp["Amount"].abs().sum(),
                    "E-mail draft": email_draft(cfg, name, key, items, cutoff)})
    requests = pd.DataFrame(req).sort_values("Amount", ascending=False) if req else pd.DataFrame(
        columns=["PAN/TAN", "Party", "Area", "Owner", "E-mail", "Points", "Amount", "E-mail draft"])

    with pd.ExcelWriter(path, engine="openpyxl") as w:
        _readme(w, ctx)
        _dashboard(w, ctx, keys, actions)
        _write(w, "Action Items", actions, "Action items - what each owner must do",
               "Sorted by priority and amount. Status/remarks come from 91_Tracker/Action_Tracker.xlsx - edit them "
               "there, not here. Amount sign: + books higher than 26AS, - books lower.",
               money=["Amount"], total_cols=["Amount"], widths={"Action": 80, "Evidence": 30})
        _write(w, "Reco PAN x FY", reco, f"26AS vs Books - TDS receivable, up to {cutoff:%d-%b-%Y}",
               f"Variance = Books - 26AS. Tolerance +/- {tol:,.0f} per party (ITD). Same layout as the manual file.",
               money=money_reco, dates=["Last party ledger received"], total_cols=money_reco,
               widths={"Action": 70, "Name": 34, "Name as per 26AS": 34})
        if len(rc.journal):
            _write(w, "Journal Entries", rc.journal, "Proposed journal entries (our books)",
                   "One JE per party. It aligns every FY to 26AS; the balancing line goes to the party account. "
                   "Review before posting. Accounts are the TDS GL most used for that FY/area.",
                   money=["Debit", "Credit"], total_cols=["Debit", "Credit"])
        _write(w, "Deductor Requests", requests, "Requests to deductors - one e-mail per party",
               "Copy the draft into mail. The e-mail column fills from 03_Masters/Party_Master.xlsx > Parties.",
               money=["Amount"], widths={"E-mail draft": 90})
        _ldc_sheets(w, ctx)
        nib = rc.not_in_books()
        _write(w, "Not in Books", nib.rename(columns={"fy": "FY", "deductor_name": "Deductor", "party_key": "PAN/TAN",
                                                     "tan": "TAN", "amount_paid": "Amount paid", "tds": "TDS"}),
               "26AS credits with nothing in books", money=["Amount paid", "TDS"], total_cols=["TDS"])
        summ, lines = rc.unallocated()
        if len(summ):
            _write(w, "Unallocated Books", lines.head(50000).rename(columns={"party_key": "Tag in books"}),
                   "Books lines without a deductor PAN",
                   "Tag each line to its PAN. 'suggested_pan_by_amount' = the amount exactly matches a 26AS line or "
                   "FY gap of a party that is short/not in books.", money=["amount"], dates=["posting_date"],
                   total_cols=["amount"])
        tim = keys[keys["who"] == WHO_NONE][["party_key", "name", "area", "books", "s26", "variance",
                                             "sub_books_pre", "sub_books_other", "sub_26as", "pending_books",
                                             "category", "action"]]
        _write(w, "Timing", tim, "Timing differences - no entry, re-check next run",
               money=["books", "s26", "variance", "sub_books_pre", "sub_books_other", "sub_26as", "pending_books"],
               total_cols=["variance"])
        _evidence_sheets(w, ctx)
        _dq_sheets(w, ctx)
        _movement_sheet(w, ctx)
    return actions


def _readme(w, ctx):
    rc, cfg = ctx["reco"], ctx["cfg"]
    rows = [
        ("Run date", dt.datetime.now().strftime("%d-%b-%Y %H:%M")),
        ("Reco cut-off", f"{rc.cutoff:%d-%b-%Y} ({ctx['cutoff_source']})"),
        ("FYs in scope", ", ".join(fy_label(f) for f in rc.fys)),
        ("Tolerance per party (Rs)", f"{cfg['tolerance']:,.0f}"),
        ("26AS files used", "; ".join(f"{fy_label(r['fy'])}: {r['source_file']} ({r['snapshot_date']})"
                                      for r in ctx["snaps"][ctx["snaps"]["used"]].to_dict("records"))
         if len(ctx["snaps"]) else "none"),
        ("GL files", "; ".join(ctx["gl_files"])),
        ("GL duplicate lines dropped", f"{ctx['gl_dupes']:,}"),
        ("Party ledgers read", str(len(ctx["ledger_info"]))),
        ("Form 16A lines read", str(len(ctx["f16a"]))),
        ("LDC certificates", str(len(ctx["ldc_master"]))),
        ("", ""),
        ("HOW TO READ", ""),
        ("Action Items", "Start here. One row per thing someone must do, with owner, amount and priority."),
        ("Reco PAN x FY", "Books vs 26AS per party and FY - the manual file's layout, with a category and action."),
        ("Journal Entries", "Entries WE pass in books (book missing TDS, reverse excess, FY reclass)."),
        ("Deductor Requests", "What each DEDUCTOR must fix, with a ready e-mail."),
        ("LDC Cut", "Cut 1: deductions above the LDC rate = cash blocked until refund (working-capital loss)."),
        ("LDC Discovery", "LDC-like rates seen in 26AS and suggested master rows, for when the LDC master is "
                          "missing/incomplete. Indicative only."),
        ("Timing", "Differences expected to clear by themselves. No entry."),
        ("Unallocated Books", "Books lines with no PAN (Contra, Other receivables...). Tag them."),
        ("Data Quality", "TANs without a PAN, 26AS files used, GL accounts, lines outside scope."),
        ("Movement", "Change vs the previous run by area and status (the manual 'Status' tab)."),
        ("", ""),
        ("WHO ACTS", ""),
        (WHO_US, "Our books are wrong or incomplete: pass the proposed JE."),
        (WHO_TAG, "Books entry has no PAN: tag it (reclass within the same GL)."),
        (WHO_DED, "Deductor must file/revise its TDS return, fix a challan, issue Form 16A or apply the LDC rate."),
        (WHO_TAX, "Tax team: LDC renewals/enhancements, write-off approvals."),
        (WHO_NONE, "Genuine timing difference. Re-check next run."),
    ]
    df = pd.DataFrame(rows, columns=["Item", "Detail"])
    df.to_excel(w, sheet_name="README", index=False, startrow=2)
    ws = w.sheets["README"]
    ws["A1"] = "26AS vs Books - automated reconciliation"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color="1F3864")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120
    for r in ws.iter_rows(min_row=3, max_row=3 + len(df)):
        for c in r:
            c.font = Font(name=FONT, size=9, bold=c.column == 1 or clean(c.value).isupper())
            c.alignment = Alignment(wrap_text=True, vertical="top")


def _dashboard(w, ctx, keys, actions):
    rc, cfg = ctx["reco"], ctx["cfg"]
    ws = w.book.create_sheet("Dashboard", 1)
    ws["A1"] = f"26AS vs Books dashboard - cut-off {rc.cutoff:%d-%b-%Y}"
    ws["A1"].font = Font(name=FONT, bold=True, size=14, color="1F3864")
    ws["A2"] = "Amounts in Rs. Variance = Books - 26AS (+ excess in books, - short in books)."
    ws["A2"].font = Font(name=FONT, italic=True, size=9)
    row = 4
    fy = rc.fyv.groupby("fy")[["books", "s26", "variance", "s26_uo"]].sum().reset_index()
    fy.loc[len(fy)] = ["Total", fy["books"].sum(), fy["s26"].sum(), fy["variance"].sum(), fy["s26_uo"].sum()]
    fy["fy"] = fy["fy"].map(lambda f: fy_label(f) if f != "Total" else "Total")
    fy.columns = ["FY", "Books", "26AS", "Variance", "of which 26AS U/O"]
    row = _block(ws, row, "1. Books vs 26AS by FY", fy, money=fy.columns[1:])

    k = keys.copy()
    k["bucket"] = np.select([k["family"] == "TAXADJ", k["who"] == WHO_OK, k["who"] == WHO_NONE],
                            ["Tax adjustments (refund/set-off)", "Matched", "Timing"], "Actionable")
    b = k.groupby("bucket")["variance"].agg(["sum", "size"]).reindex(
        ["Matched", "Timing", "Tax adjustments (refund/set-off)", "Actionable"]).fillna(0)
    b = b.reset_index()
    b.columns = ["Bucket", "Variance", "Parties"]
    b.loc[len(b)] = ["Total", b["Variance"].sum(), b["Parties"].sum()]
    row = _block(ws, row, "2. Where the variance sits", b, money=["Variance", "Parties"])

    k = k[k["family"] != "TAXADJ"]
    st = k.pivot_table(index="status", columns="area", values="variance", aggfunc="sum", fill_value=0) / 1e6
    st["Total"] = st.sum(axis=1)
    st = st.reset_index().rename(columns={"status": "Status (Rs mn)"})
    row = _block(ws, row, "3. Status by area (Rs mn), excluding tax adjustments - as in the manual 'Status' tab", st,
                 money=[],
                 pct=[])
    for r in ws.iter_rows(min_row=row - len(st) - 2, max_row=row - 3, min_col=2, max_col=len(st.columns)):
        for c in r:
            c.number_format = '#,##0.00;(#,##0.00);"-"'
    stc = k.pivot_table(index="status", columns="area", values="variance", aggfunc="size", fill_value=0)
    stc["Total"] = stc.sum(axis=1)
    stc = stc.reset_index().rename(columns={"status": "Status (count)"})
    row = _block(ws, row, "4. Status by area (count of parties)", stc)

    a = actions.copy()
    a["abs"] = a["Amount"].abs()
    wa = a.groupby(["Who Acts", "Category"]).agg(Items=("Action ID", "size"), Amount=("abs", "sum")).reset_index()
    wa = wa.sort_values(["Who Acts", "Amount"], ascending=[True, False])
    wa.loc[len(wa)] = ["Total", "", wa["Items"].sum(), wa["Amount"].sum()]
    row = _block(ws, row, "5. Open actions by who acts (absolute amounts)", wa, money=["Amount", "Items"])

    ls = ctx["ldc_summary"]
    if len(ls):
        l = ls.groupby("fy").agg(TDS=("tds_deposited", "sum"), Expected=("expected_tds", "sum"),
                                 Excess=("excess", "sum"), WC=("wc_cost_net", "sum")).reset_index()
        l["fy"] = l["fy"].map(fy_label)
        l.columns = ["FY", "TDS on LDC-covered payments", "Expected at LDC rate", "Excess (blocked cash)",
                     "Working-capital cost (net)"]
        row = _block(ws, row, "6. LDC cut - excess deduction and its cost", l, money=l.columns[1:])
    else:
        ind = ctx["ldc_indicative"]
        if len(ind):
            l = ind.groupby("fy").agg(Parties=("tan", "nunique"), TDS=("tds_deposited", "sum"),
                                      Excess=("indicative_excess", "sum")).reset_index()
            l["fy"] = l["fy"].map(fy_label)
            l.columns = ["FY", "Deductors", "TDS deducted at the higher rate",
                         "Indicative excess vs the deductor's own LDC rate"]
            row = _block(ws, row, "6. LDC cut - NO LDC MASTER YET. Indicative: deductors who applied an LDC-like rate "
                                  "on some payments and the standard rate on others (see LDC Discovery)", l,
                         money=l.columns[1:])
    mv = ctx.get("movement")
    if mv is not None and len(mv) and "amount_prev" in mv:
        m = mv.copy()
        m[["amount_now", "amount_prev", "amount_change"]] /= 1e6
        m = m[["area", "status", "amount_prev", "amount_now", "amount_change", "count_prev", "count_now",
               "count_change"]]
        m.columns = ["Area", "Status", "Prev (Rs mn)", "Now (Rs mn)", "Change (Rs mn)", "Prev #", "Now #",
                     "Change #"]
        row = _block(ws, row, "7. Movement vs previous run", m, money=["Prev #", "Now #", "Change #"])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 44
    for col in range(3, 16):
        ws.column_dimensions[get_column_letter(col)].width = 17


def _ldc_sheets(w, ctx):
    ls, lines, util = ctx["ldc_summary"], ctx["ldc_lines"], ctx["ldc_util"]
    if len(ls):
        d = ls.copy()
        d["fy"] = d["fy"].map(fy_label)
        d["ldc_rate"] = d["ldc_rate"]
        d = d.sort_values("excess", ascending=False)
        _write(w, "LDC Cut", d, "Cut 1 - LDC rate vs rate applied by the deductor",
               "excess = TDS deposited - (amount x LDC rate) on payments within the certificate's limit and validity. "
               "wc_cost_net = excess x cost of capital x days blocked - refund interest (Settings).",
               money=["amount_paid", "tds_deposited", "expected_tds", "excess", "short", "wc_cost_net"],
               pct=["ldc_rate"], dates=["first_above", "last_txn"], total_cols=["excess", "wc_cost_net"])
    else:
        pd.DataFrame({"Note": ["No LDC master rows yet. Fill 03_Masters/LDC_Certificates.xlsx (one row per "
                               "certificate: TAN, section, rate %, validity, amount limit) and re-run. "
                               "'LDC Discovery' lists what 26AS suggests."]}) \
            .to_excel(w, sheet_name="LDC Cut", index=False)
    if len(util):
        _write(w, "LDC Certificates", util, "LDC certificate utilisation",
               money=["Amount Limit (Rs)", "Amount Consumed (Rs)"], pct=["% Used"], dates=["Valid From", "Valid To"])
    if len(lines):
        cols = ["flag", "party_key", "tan", "deductor_name", "section_code", "section", "fy", "txn_date",
                "booking_status", "amount_paid", "tds_deposited", "rate", "std_rate", "cert", "ldc_rate",
                "expected_tds", "excess_over_ldc", "short_vs_ldc", "excess_over_std", "blocked_days", "wc_cost_net"]
        _write(w, "LDC Lines", lines[cols].head(100000), "26AS lines with a rate exception",
               money=["amount_paid", "tds_deposited", "expected_tds", "excess_over_ldc", "short_vs_ldc",
                      "excess_over_std", "wc_cost_net"], pct=["rate", "std_rate", "ldc_rate"], dates=["txn_date"],
               total_cols=["excess_over_ldc", "excess_over_std", "wc_cost_net"])
    prof, ind, sug = ctx["ldc_profile"], ctx["ldc_indicative"], ctx["ldc_suggest"]
    if len(ind):
        _write(w, "LDC Discovery", ind, "Deductors who applied an LDC-like rate on some payments and the standard "
                                        "rate on others (same section, same FY)",
               "INDICATIVE until the LDC master is filled. The benchmark is the lowest rate THIS deductor applied in the "
               "FY, so the certificate very likely exists. Confirm it, add it to the master, and chase the deductor.",
               money=["amount_paid", "tds_deposited", "indicative_excess"], pct=["applied_rate", "probable_ldc_rate"],
               total_cols=["indicative_excess"])
    if len(sug):
        _write(w, "LDC Master Suggestions", sug, "Sub-standard rates seen in 26AS, not yet in the LDC master",
               "Check each against the actual certificate, then copy it to 03_Masters/LDC_Certificates.xlsx.",
               money=["amount_paid"], pct=["observed_rate"], dates=["from_date", "to_date"])
    if len(prof):
        _write(w, "Rate Profile", prof, "Rate profile by section and quarter", money=["amount_paid", "tds"],
               pct=["r", "std_rate", "share_of_lines", "probable_ldc_rate"])


def _evidence_sheets(w, ctx):
    li, ll, f16, fcmp = ctx["ledger_info"], ctx["ledger_lines"], ctx["f16a"], ctx["f16a_cmp"]
    if len(li):
        _write(w, "Party Ledgers", li, "Party ledgers read from 04_Party_Ledgers",
               "Put the PAN or TAN in the file name (e.g. AAJCG0980D_ledger_FY25.xlsx) for a sure match.",
               money=["closing_balance"], dates=["from", "to", "received_on"])
    if len(ll):
        _write(w, "Ledger TDS Lines", ll, "TDS lines picked from party ledgers", money=["tds"], dates=["date"],
               total_cols=["tds"])
    if len(f16):
        _write(w, "Form 16A", f16, "TDS certificates read from 05_Form16A",
               money=["amount_paid", "tax_deducted", "tax_deposited"])
    if len(fcmp):
        _write(w, "16A vs 26AS", fcmp, "Form 16A vs 26AS by TAN and quarter",
               money=["form16a_tds", "tds_26as", "difference"])


def _dq_sheets(w, ctx):
    rc = ctx["reco"]
    sug = ctx["tan_suggestions"]
    if len(sug):
        _write(w, "Unmapped TANs", sug, "26AS TANs without a confirmed PAN",
               "Reco uses the TAN as the key until you confirm a PAN in 03_Masters/Party_Master.xlsx > TAN_Map "
               "(set Status = Confirmed). New TANs are appended there automatically.", money=["TDS"])
    if len(ctx["snaps"]):
        s = ctx["snaps"].copy()
        s["fy"] = s["fy"].map(fy_label)
        _write(w, "26AS Snapshots", s, "26AS files found - newest per FY is used", money=["tds_deposited"],
               dates=["last_booking"])
    if len(ctx["gl_accounts"]):
        _write(w, "GL Accounts", ctx["gl_accounts"], "GL accounts in the books data", money=["Balance"])
    out = rc.books_out
    if len(out):
        o = out.groupby(["fy_col", "gl_code"], dropna=False)["amount"].agg(["size", "sum"]).reset_index()
        o.columns = ["Fiscal (as in GL)", "GL", "Lines", "Amount"]
        _write(w, "Out of Scope Books", o, "Books lines outside the FY range (e.g. opening balances)",
               money=["Amount"], total_cols=["Amount"])


def _movement_sheet(w, ctx):
    pk = ctx.get("movement_keys")
    if pk is not None and len(pk):
        _write(w, "Movement", pk, f"Change vs previous run ({ctx.get('prev_meta', {}).get('run_at', '')})",
               money=["variance", "prev_variance", "change"], total_cols=["change"])
